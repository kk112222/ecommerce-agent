"""文档解析器 —— PDF / Word / 网页 / Markdown → 纯文本"""
from pathlib import Path


def parse_file(filepath: Path) -> str:
    """根据文件后缀自动选解析器，返回纯文本"""
    suffix = filepath.suffix.lower()

    if suffix == ".md":
        return _parse_md(filepath)
    elif suffix == ".pdf":
        return _parse_pdf(filepath)
    elif suffix == ".docx":
        return _parse_docx(filepath)
    elif suffix in (".html", ".htm"):
        return _parse_html(filepath)
    elif suffix in (".csv", ".tsv"):
        return _parse_csv(filepath)
    else:
        # 未知格式当纯文本读
        return filepath.read_text(encoding="utf-8", errors="replace")


def _parse_md(filepath: Path) -> str:
    return filepath.read_text(encoding="utf-8")


def _parse_csv(filepath: Path) -> str:
    """CSV/TSV → 表格文本（标准库 csv，零依赖）

    utf-8-sig 兼容 Excel 导出时带 BOM 头的情况（否则第一列会多个 ﻿）
    """
    import csv
    delimiter = "," if filepath.suffix.lower() == ".csv" else "\t"
    with open(filepath, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=delimiter))
    # 每行用 | 拼列，方便 LLM 读表格结构
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _parse_pdf(filepath: Path) -> str:
    """PDF → 提取每一页文字"""
    import fitz  # pymupdf
    doc = fitz.open(str(filepath))
    pages = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            pages.append(text.strip())
    doc.close()
    return "\n\n".join(pages)


def _parse_docx(filepath: Path) -> str:
    """Word .docx → 提取段落文字"""
    from docx import Document
    doc = Document(str(filepath))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # 也读表格里的文字
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n\n".join(paragraphs)


def _parse_html(filepath: Path) -> str:
    """网页 → 提取正文 → Markdown 文本"""
    from bs4 import BeautifulSoup
    import html2text

    html = filepath.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html, "lxml")

    # 去掉 script/style/nav/footer 等噪音
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    # 尝试找正文区域
    main = soup.find("main") or soup.find("article") or soup.find("body")
    if main is None:
        return ""

    h = html2text.HTML2Text()
    h.ignore_links = False
    h.body_width = 0  # 不自动换行
    return h.handle(str(main))
